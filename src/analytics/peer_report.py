"""
peer_report.py
--------------
Generates output/peer_comparison.xlsx: one sheet per peer group (11
sheets), with each company's metric values plus percentile rank,
colour-coded, benchmark company highlighted, and a peer-group median
summary row.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.analytics.peer import (
    PEER_METRICS,
    compute_peer_percentiles,
    load_peer_groups,
)
from src.screener.engine import load_screener_data

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "output" / "peer_comparison.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BENCHMARK_FILL = PatternFill(
    start_color="FFD966", end_color="FFD966", fill_type="solid"
)  # gold/amber
MEDIAN_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _percentile_fill(pct):
    if pct is None or (isinstance(pct, float) and pct != pct):
        return None
    if pct >= 75:
        return GREEN_FILL
    if pct >= 25:
        return YELLOW_FILL
    return RED_FILL


def _sheet_name(name):
    return str(name)[:31]


def generate_peer_comparison_report(db_path, output_path=None):
    """Generate peer comparison report."""
    output_path = Path(output_path) if output_path else DEFAULT_OUTPUT_PATH
    output_path.parent.mkdir(parents=True, exist_ok=True)

    metrics_df = load_screener_data(db_path)
    peer_groups = load_peer_groups(db_path)
    percentiles = compute_peer_percentiles(db_path)

    merged = peer_groups.merge(
        metrics_df[["company_id", "company_name"] + [c for c, _ in PEER_METRICS.values()]],
        on="company_id",
        how="left",
    )

    metric_names = list(PEER_METRICS.keys())
    metric_cols = {name: col for name, (col, _) in PEER_METRICS.items()}

    wb = Workbook()
    wb.remove(wb.active)

    for group_name, group in merged.groupby("peer_group_name"):

        ws = wb.create_sheet(_sheet_name(group_name))

        # ---------------- Header ----------------
        headers = ["company_id", "company_name"]
        for metric in metric_names:
            headers.append(metric)
            headers.append(f"{metric} %ile")

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        group_percentiles = percentiles[percentiles["peer_group_name"] == group_name]

        # ---------------- Data rows ----------------
        row_idx = 2
        for _, comp_row in group.iterrows():

            ws.cell(row=row_idx, column=1, value=comp_row["company_id"])
            ws.cell(row=row_idx, column=2, value=comp_row["company_name"])

            col_idx = 3
            for metric in metric_names:
                value = comp_row.get(metric_cols[metric])
                if isinstance(value, float) and value != value:
                    value = None
                ws.cell(row=row_idx, column=col_idx, value=value)

                pct_row = group_percentiles[
                    (group_percentiles["company_id"] == comp_row["company_id"])
                    & (group_percentiles["metric"] == metric)
                ]
                pct = pct_row["percentile_rank"].iloc[0] if not pct_row.empty else None

                pct_cell = ws.cell(row=row_idx, column=col_idx + 1, value=pct)
                fill = _percentile_fill(pct)
                if fill:
                    pct_cell.fill = fill

                col_idx += 2

            # Highlight the benchmark company row (gold/amber)
            if bool(comp_row.get("is_benchmark")):
                for c in range(1, len(headers) + 1):
                    existing_fill = ws.cell(row=row_idx, column=c).fill
                    if existing_fill is None or existing_fill.fgColor.rgb in (None, "00000000"):
                        ws.cell(row=row_idx, column=c).fill = BENCHMARK_FILL

            row_idx += 1

        # ---------------- Median summary row ----------------
        median_row = row_idx
        ws.cell(row=median_row, column=1, value="Peer Group Median").font = Font(bold=True)
        for c in range(1, len(headers) + 1):
            ws.cell(row=median_row, column=c).fill = MEDIAN_FILL

        col_idx = 3
        for metric in metric_names:
            values = group[metric_cols[metric]].dropna()
            median_value = float(values.median()) if not values.empty else None
            ws.cell(row=median_row, column=col_idx, value=median_value)

            pct_values = group_percentiles[group_percentiles["metric"] == metric][
                "percentile_rank"
            ].dropna()
            median_pct = float(pct_values.median()) if not pct_values.empty else None
            ws.cell(row=median_row, column=col_idx + 1, value=median_pct)

            col_idx += 2

        # ---------------- Column widths ----------------
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 16

        ws.freeze_panes = "C2"

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    path = generate_peer_comparison_report(PROJECT_ROOT / "data" / "nifty100.db")
    print(f"Wrote {path}")
