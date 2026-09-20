"""
report.py
---------
Generates output/screener_output.xlsx: one sheet per preset defined in
screener_config.yaml, with KPI columns and threshold-based colour
coding (green = meets the preset's own threshold for that metric,
red = fails it).
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.screener.engine import (
    FILTER_MAP,
    load_config,
    run_all_presets,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "output" / "screener_output.xlsx"

# The ~20 KPI columns shown on every preset sheet.
KPI_COLUMNS = [
    "company_id",
    "company_name",
    "broad_sector",
    "return_on_equity_pct",
    "roce_percentage",
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "debt_to_equity",
    "interest_coverage",
    "asset_turnover",
    "free_cash_flow_cr",
    "cfo_pat_ratio",
    "dividend_payout_ratio_pct",
    "dividend_yield_pct",
    "pe_ratio",
    "pb_ratio",
    "market_cap_crore",
    "sales",
    "compounded_sales_growth",
    "compounded_profit_growth",
    "composite_quality_score",
    "sector_rank",
    "overall_rank",
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _meets_threshold(rule, column, value, threshold):
    """
    Returns True/False/None (None = not evaluable, e.g. missing data
    or a rule with no direct column, so the cell is left uncoloured).
    """
    if value is None or (isinstance(value, float) and value != value):  # NaN check
        # icr_min: NaN means debt-free, which always passes
        if rule == "icr_min":
            return True
        return None

    if rule == "de_declining_yoy":
        return bool(value) == bool(threshold)

    if rule.endswith("_min"):
        return value >= threshold

    if rule.endswith("_max"):
        return value <= threshold

    return None


def _sheet_name(preset_name):
    # Excel sheet names: max 31 chars, no special chars
    return preset_name.replace("_", " ").title()[:31]


def generate_screener_report(db_path, output_path=None):
    """
    Build output/screener_output.xlsx with one sheet per preset.
    Returns the path written to.
    """
    output_path = Path(output_path) if output_path else DEFAULT_OUTPUT_PATH
    output_path.parent.mkdir(parents=True, exist_ok=True)

    config = load_config()
    results = run_all_presets(db_path)

    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    for preset_name, df in results.items():

        rules = config.get(preset_name, {})
        df = df.sort_values("composite_quality_score", ascending=False).reset_index(drop=True)

        ws = wb.create_sheet(_sheet_name(preset_name))

        # Header row
        for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        # Map: column -> rule (for threshold colouring), since a
        # column can be governed by at most one rule per preset here.
        column_to_rule = {}
        for rule in rules:
            col = FILTER_MAP.get(rule)
            if col:
                column_to_rule[col] = rule

        # Data rows
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            row_dict = row._asdict()

            for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
                value = row_dict.get(col_name)

                # pandas NA / NaN -> blank cell
                if value is not None and isinstance(value, float) and value != value:
                    value = None

                cell = ws.cell(row=row_idx, column=col_idx, value=value)

                rule = column_to_rule.get(col_name)
                if rule is not None:
                    threshold = rules[rule]
                    meets = _meets_threshold(rule, col_name, value, threshold)
                    if meets is True:
                        cell.fill = GREEN_FILL
                    elif meets is False:
                        cell.fill = RED_FILL

            # de_declining_yoy isn't a KPI column, but still colour
            # the debt_to_equity cell for it if that's the active rule
            if "de_declining_yoy" in rules and "debt_to_equity" in KPI_COLUMNS:
                de_col_idx = KPI_COLUMNS.index("debt_to_equity") + 1
                de_declining = row_dict.get("de_declining_yoy")
                cell = ws.cell(row=row_idx, column=de_col_idx)
                if de_declining is True:
                    cell.fill = GREEN_FILL
                elif de_declining is False:
                    cell.fill = RED_FILL

        # Column widths
        for col_idx, col_name in enumerate(KPI_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(col_name) + 2)

        ws.freeze_panes = "A2"

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    path = generate_screener_report(PROJECT_ROOT / "data" / "nifty100.db")
    print(f"Wrote {path}")
